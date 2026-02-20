"""PowerPoint generation for hymn plans and service programs."""

import re
from io import BytesIO
from typing import Optional

from sqlalchemy.orm import Session

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor

from utils import clean_text, get_slides_by_type

BG_COLOR = RGBColor(0, 0, 0)
TEXT_COLOR = RGBColor(255, 255, 255)


class PptxConfig:
    """Font / size settings for PPTX generation."""

    def __init__(
        self,
        font_select: str = "Arial",
        font_manual: str = "",
        title_size: int = 80,
        lyrics_size: int = 60,
    ):
        if font_select == "Manual" and font_manual.strip():
            self.font_name = font_manual.strip()
        else:
            self.font_name = font_select
        self.title_size = Pt(title_size)
        self.number_size = Pt(int(title_size * 0.5))
        self.lyrics_size = Pt(lyrics_size)
        self._title_size_raw = title_size
        self._lyrics_size_raw = lyrics_size


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _geometry(aspect: str):
    """Return (slide_width, slide_height, margin) for the given aspect ratio."""
    if aspect == "4:3":
        w, h = Inches(10), Inches(7.5)
    else:
        w, h = Inches(16), Inches(9)
    margin = Inches(0.5)
    return w, h, margin


def _create_black_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    return slide


def _add_text_frame(slide, margin, safe_w, safe_h):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, margin, margin, safe_w, safe_h
    )
    shape.fill.background()
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tf


def _duplicate_slide(pres, index):
    """Duplicate the slide at *index* to the end of the presentation."""
    source = pres.slides[index]
    dest = pres.slides.add_slide(source.slide_layout)

    # Copy background
    try:
        if source.background.fill.type == 1:
            dest.background.fill.solid()
            dest.background.fill.fore_color.rgb = source.background.fill.fore_color.rgb
    except:
        pass

    # Copy shapes
    for shape in source.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE or shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                    geom = MSO_SHAPE.RECTANGLE
                else:
                    geom = shape.auto_shape_type

                new_shape = dest.shapes.add_shape(
                    geom, shape.left, shape.top, shape.width, shape.height
                )

                try:
                    if shape.fill.type == 1:
                        new_shape.fill.solid()
                        new_shape.fill.fore_color.rgb = shape.fill.fore_color.rgb
                except:
                    pass

                if shape.has_text_frame:
                    new_shape.text_frame.clear()
                    for paragraph in shape.text_frame.paragraphs:
                        new_p = new_shape.text_frame.add_paragraph()
                        new_p.text = paragraph.text
                        new_p.alignment = paragraph.alignment
                        if paragraph.runs:
                            r_source = paragraph.runs[0]
                            r_dest = new_p.font
                            r_dest.name = r_source.font.name
                            r_dest.size = r_source.font.size
                            r_dest.bold = r_source.font.bold
                            r_dest.italic = r_source.font.italic
                            r_dest.underline = r_source.font.underline
                            try:
                                if r_source.font.color.type == 1:
                                    r_dest.color.rgb = r_source.font.color.rgb
                            except:
                                pass
            except Exception as e:
                print(f"Skipping shape due to error: {e}")

        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                blob = shape.image.blob
                dest.shapes.add_picture(
                    BytesIO(blob), shape.left, shape.top, shape.width, shape.height
                )
            except:
                pass

    return dest


def _move_slide(pres, old_index, new_index):
    xml_slides = pres.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])


def _delete_slides_with_pattern(prs, pattern):
    """Delete all slides containing the given text pattern."""
    slides_to_delete_indices = []

    for i, slide in enumerate(prs.slides):
        found_pattern = False
        for shape in slide.shapes:
            if shape.has_text_frame:
                if shape.text and pattern in shape.text:
                    found_pattern = True
                    break
        if found_pattern:
            slides_to_delete_indices.append(i)

    if slides_to_delete_indices:
        xml_slides = prs.slides._sldIdLst
        slides_element_list = list(xml_slides)
        for index in sorted(slides_to_delete_indices, reverse=True):
            xml_slides.remove(slides_element_list[index])

    return prs


# ---------------------------------------------------------------------------
# Public generators
# ---------------------------------------------------------------------------

def generate_hymn_plan_pptx(plan_items, aspect: str = "4:3", config: Optional[PptxConfig] = None) -> BytesIO:
    """Generate a black-slide PPTX from a list of ServicePlanHymn items."""
    cfg = config or PptxConfig()
    prs = Presentation()

    slide_w, slide_h, margin = _geometry(aspect)
    prs.slide_width = slide_w
    prs.slide_height = slide_h
    safe_w = slide_w - margin * 2
    safe_h = slide_h - margin * 2

    for index, item in enumerate(plan_items):
        hymn = item.hymn
        if not hymn:
            continue

        # Title slide
        slide = _create_black_slide(prs)
        tf = _add_text_frame(slide, margin, safe_w, safe_h)

        p1 = tf.paragraphs[0]
        p1.text = f"Hymn #{hymn.number}"
        p1.font.name = cfg.font_name
        p1.font.size = cfg.number_size
        p1.font.bold = False
        p1.font.color.rgb = TEXT_COLOR
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = hymn.title
        p2.font.name = cfg.font_name
        p2.font.size = cfg.title_size
        p2.font.bold = True
        p2.font.color.rgb = TEXT_COLOR
        p2.alignment = PP_ALIGN.CENTER

        # Lyrics slides
        slides_source = get_slides_by_type(hymn, preferred_type="PPT")

        grouped = []
        current_group = None
        for s in slides_source:
            content = clean_text(s.content)
            if s.label or current_group is None:
                if current_group:
                    grouped.append(current_group)
                current_group = {"content": content}
            else:
                current_group["content"] += "\n\n" + content
        if current_group:
            grouped.append(current_group)

        for g in grouped:
            slide = _create_black_slide(prs)
            tf = _add_text_frame(slide, margin, safe_w, safe_h)
            p = tf.paragraphs[0]
            p.text = g["content"]
            p.font.name = cfg.font_name
            p.font.size = cfg.lyrics_size
            p.font.bold = False
            p.font.color.rgb = TEXT_COLOR
            p.alignment = PP_ALIGN.CENTER

        # Spacer between hymns
        if index < len(plan_items) - 1:
            _create_black_slide(prs)

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output


def generate_program_pptx(
    program,
    db: Session,
    template_path: str,
    config: Optional[PptxConfig] = None,
) -> BytesIO:
    """Generate a template-based PPTX with placeholder replacement and hymn explosion."""
    from models import ServicePlanHymnModel

    cfg = config or PptxConfig()
    prs = Presentation(template_path)

    # --- PHASE 1: STANDARD REPLACEMENTS ---
    replacements = {}
    for item in program.items:
        if item.tag and item.subtitle:
            replacements[item.tag] = item.subtitle

    search_replace_pairs = {("{{" + key + "}}"): str(value) for key, value in replacements.items()}

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    full_text = paragraph.text
                    for pattern, replacement in search_replace_pairs.items():
                        if pattern in full_text:
                            full_text = full_text.replace(pattern, replacement)
                            paragraph.text = full_text

    # --- PHASE 2: HYMN EXPLOSION ---
    # Iterate backwards to maintain valid indices for upcoming slides
    for i in range(len(prs.slides) - 1, -1, -1):
        slide = prs.slides[i]

        # 1. Check for Hymn Tag
        hymn_seq = None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                match = re.search(r"\{\{hymn_(\d+)\}\}", paragraph.text)
                if match:
                    hymn_seq = int(match.group(1))
                    break
            if hymn_seq:
                break

        if not hymn_seq:
            continue

        # 2. Fetch Content
        plan_item = db.query(ServicePlanHymnModel).filter(
            ServicePlanHymnModel.sequence == hymn_seq
        ).first()
        content_chunks = []

        if not plan_item or not plan_item.hymn:
            content_chunks.append(f"(Hymn #{hymn_seq} not scheduled)")
        else:
            hymn = plan_item.hymn
            content_chunks.append(f"Hymn #{hymn.number}\n{hymn.title}")
            ppt_slides = [s for s in hymn.slides if s.type == "PPT"]
            if ppt_slides:
                for s in sorted(ppt_slides, key=lambda x: x.order):
                    content_chunks.append(s.content)

        # 3. Apply Content
        insertion_index = i

        for idx, extra_chunk in enumerate(content_chunks):
            new_slide = _duplicate_slide(prs, i)
            _move_slide(prs, len(prs.slides) - 1, insertion_index + 1)

            if idx == 0:
                for s in new_slide.shapes:
                    if s.has_text_frame:
                        p_element = s.text_frame.paragraphs[0]._p
                        p_element.getparent().remove(p_element)

                        p1 = s.text_frame.add_paragraph()
                        p1.text = extra_chunk.splitlines()[0].strip()
                        p1.font.name = cfg.font_name
                        p1.font.size = Pt(int(cfg._title_size_raw * 0.5))
                        p1.font.bold = False
                        p1.alignment = PP_ALIGN.CENTER

                        p2 = s.text_frame.add_paragraph()
                        p2.text = "\n".join(extra_chunk.splitlines()[1:]).strip()
                        p2.font.name = cfg.font_name
                        p2.font.size = Pt(cfg._title_size_raw)
                        p2.font.bold = True
                        p2.alignment = PP_ALIGN.CENTER
            else:
                for s in new_slide.shapes:
                    if s.has_text_frame:
                        for p in s.text_frame.paragraphs:
                            p.text = p.text.replace(p.text, clean_text(extra_chunk))
                            p.font.name = cfg.font_name
                            p.font.size = Pt(cfg._lyrics_size_raw)

            insertion_index += 1

    _delete_slides_with_pattern(prs, "{{hymn_")

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output
