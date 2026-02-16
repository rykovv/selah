from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import os
import copy

def copy_run_format(source_run, target_run):
    """
    Copy font styles (bold, italic, underline, size, color, font name) 
    from source_run to target_run.
    """
    target_run.font.bold = source_run.font.bold
    target_run.font.italic = source_run.font.italic
    target_run.font.underline = source_run.font.underline
    
    if source_run.font.size is not None:
        target_run.font.size = source_run.font.size
    
    if source_run.font.name is not None:
        target_run.font.name = source_run.font.name
        
    # Copy color if it is defined
    try:
        if source_run.font.color.type == 1: # RGB
            target_run.font.color.rgb = source_run.font.color.rgb
        elif source_run.font.color.type == 2: # Theme color
            target_run.font.color.theme_color = source_run.font.color.theme_color
            # Brightness/LumMod is tricky to copy perfectly, but this covers basics
    except ValueError:
        pass # Color might not be set or complex

def replace_text_preserving_formatting(pptx_path, output_path, replacements):
    if not os.path.exists(pptx_path):
        print(f"Error: File '{pptx_path}' not found.")
        return

    prs = Presentation(pptx_path)
    
    # Prepare keys with double braces
    search_replace_pairs = {f"{{{{{key}}}}}": value for key, value in replacements.items()}
    
    replacements_count = 0

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            
            for paragraph in shape.text_frame.paragraphs:
                # 1. Check if any pattern exists in the FULL paragraph text
                # We do this before touching runs to avoid breaking things unnecessarily
                full_text = paragraph.text
                match_found = False
                
                for pattern, replacement in search_replace_pairs.items():
                    if pattern in full_text:
                        full_text = full_text.replace(pattern, replacement)
                        match_found = True
                        replacements_count += 1
                
                if match_found:
                    # 2. Capture the style of the first run 
                    # (Assuming the whole placeholder shares the style of the start)
                    original_runs = paragraph.runs
                    if len(original_runs) > 0:
                        # We copy the style object to a temporary holder or just reference the first run
                        # Since we are about to clear the paragraph, we must grab values now
                        first_run = original_runs[0]
                        
                        # We save specific attributes because the run object itself will be lost
                        saved_font_props = {
                            'bold': first_run.font.bold,
                            'italic': first_run.font.italic,
                            'underline': first_run.font.underline,
                            'size': first_run.font.size,
                            'name': first_run.font.name,
                            'color_obj': first_run.font.color if first_run.font.color else None
                        }
                    else:
                        saved_font_props = None

                    # 3. Clear existing runs and create a new single run with the new text
                    paragraph.clear() 
                    new_run = paragraph.add_run()
                    new_run.text = full_text

                    # 4. Re-apply the saved formatting
                    if saved_font_props:
                        new_run.font.bold = saved_font_props['bold']
                        new_run.font.italic = saved_font_props['italic']
                        new_run.font.underline = saved_font_props['underline']
                        new_run.font.size = saved_font_props['size']
                        new_run.font.name = saved_font_props['name']
                        
                        # Re-apply color
                        if saved_font_props['color_obj']:
                            try:
                                if saved_font_props['color_obj'].type == 1: # RGB
                                    new_run.font.color.rgb = saved_font_props['color_obj'].rgb
                                elif saved_font_props['color_obj'].type == 2: # Theme
                                    new_run.font.color.theme_color = saved_font_props['color_obj'].theme_color
                            except (AttributeError, ValueError):
                                pass

    prs.save(output_path)
    print(f"Done! Replaced {replacements_count} patterns.")
    print(f"Saved to: {output_path}")

# --- Configuration ---

input_file = r"D:\Home\Downloads\Redding_SDA_Church_Service_Templ.pptx"
output_file = "service_presentation.pptx"

data = {
    "congregational_prayer": "Marty Harryman",
    "special_music": "Juluis Sisona",
    "scripture_reading": "Miranda Rincon",
    "message_title": "Unequally Yoked",
    "speaker_name": "Pastor Robert Fisher"
}

# --- Run ---
if __name__ == "__main__":
    replace_text_preserving_formatting(input_file, output_file, data)